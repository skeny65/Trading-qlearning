"""
excel_logger.py - Registro Excel mensual por carpeta.

- Un archivo por mes: trade_log_YYYY_MM.xlsx
- Cada fila es un trade (se escribe en OPEN y se actualiza en CLOSE)
- Incluye: balance_antes, balance_despues, profit_usdt, duracion_min, precio_salida
- Al inicio de cada mes, se agrega una fila de balance inicial heredado
"""
import logging
import os
from datetime import datetime, timezone
from pathlib import Path

import config

logger = logging.getLogger("bot3.excel")

COLUMNS = [
    "fecha",            # 2026-05-28
    "hora_utc",         # 14:32:01
    "event_id",
    "id_carpeta",
    "emisor",
    "mode",             # DRY_RUN / LIVE
    "symbol",
    "side",             # buy / sell
    "precio_entrada",
    "precio_salida",    # PENDING → precio al cerrar
    "sl",
    "tp",
    "margen_usdt",
    "leverage",
    "notional_usdt",    # margen × leverage
    "order_id",
    "binance_status",
    "duracion_min",     # PENDING → minutos hasta el cierre
    "resultado",        # PENDING → WIN / LOSS
    "pnl_pct",          # PENDING → "+1.45%"
    "profit_usdt",      # PENDING → ganancia/pérdida real en USDT
    "close_reason",
    "balance_antes",    # balance al abrir
    "balance_despues",  # balance tras cerrar (PENDING hasta el cierre)
]

_COL_WIDTHS = {
    "fecha": 12, "hora_utc": 10, "event_id": 24, "id_carpeta": 10,
    "emisor": 13, "mode": 9, "symbol": 10, "side": 6,
    "precio_entrada": 14, "precio_salida": 14, "sl": 10, "tp": 10,
    "margen_usdt": 12, "leverage": 9, "notional_usdt": 13,
    "order_id": 26, "binance_status": 14,
    "duracion_min": 12, "resultado": 9,
    "pnl_pct": 10, "profit_usdt": 12, "close_reason": 25,
    "balance_antes": 14, "balance_despues": 15,
}

# Colores por resultado para la fila completa
_COLOR_WIN    = "E2EFDA"   # verde suave
_COLOR_LOSS   = "FCE4D6"   # rojo suave
_COLOR_HEADER = "1F4E79"   # azul oscuro
_COLOR_INICIO = "FFF2CC"   # amarillo para fila de balance inicial del mes


def get_excel_path(strategy_id: str, dt: datetime = None) -> str:
    """
    Retorna la ruta del Excel del mes correspondiente.
    trade_log_YYYY_MM.xlsx dentro de la carpeta de la estrategia.
    """
    if dt is None:
        dt = datetime.now(timezone.utc)
    filename = f"trade_log_{dt.year:04d}_{dt.month:02d}.xlsx"
    return os.path.join(config.strategy_data_dir(strategy_id), filename)


def _format_header(ws) -> None:
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
        fill = PatternFill("solid", fgColor=_COLOR_HEADER)
        font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            if cell.value:
                cell.font      = font
                cell.fill      = fill
                cell.alignment = Alignment(horizontal="center")
                ws.column_dimensions[cell.column_letter].width = _COL_WIDTHS.get(
                    str(cell.value).lower(), 14
                )
        ws.freeze_panes = "A2"
    except Exception as e:
        logger.debug(f"Formato de cabecera no aplicado: {e}")


def _color_row(ws, row_num: int, resultado: str) -> None:
    """Colorea toda la fila según WIN / LOSS."""
    try:
        from openpyxl.styles import PatternFill
        if resultado == "WIN":
            color = _COLOR_WIN
        elif resultado == "LOSS":
            color = _COLOR_LOSS
        else:
            return
        fill = PatternFill("solid", fgColor=color)
        for cell in ws[row_num]:
            cell.fill = fill
    except Exception:
        pass


def _get_or_create_wb(excel_path: str, strategy_id: str, is_new_month: bool = False):
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl no instalado: pip install openpyxl")

    os.makedirs(os.path.dirname(excel_path), exist_ok=True)

    if os.path.exists(excel_path):
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        return wb, ws, False

    # Crear nuevo archivo
    wb = openpyxl.Workbook()
    ws = wb.active
    month_str = Path(excel_path).stem.replace("trade_log_", "")  # "2026_05"
    ws.title  = f"trades_{month_str}"
    ws.append(COLUMNS)
    _format_header(ws)

    # Fila de balance inicial del mes (heredado del mes anterior)
    if is_new_month:
        _add_balance_inicio_row(ws, strategy_id)

    return wb, ws, True


def _add_balance_inicio_row(ws, strategy_id: str) -> None:
    """
    Agrega una fila especial mostrando el balance heredado del mes anterior.
    """
    try:
        from utils.balance_tracker import get_balance
        from openpyxl.styles import PatternFill, Font
        balance = get_balance(strategy_id)
        now     = datetime.now(timezone.utc)

        row_data = {col: "" for col in COLUMNS}
        row_data["fecha"]           = now.strftime("%Y-%m-01")
        row_data["hora_utc"]        = "00:00:00"
        row_data["id_carpeta"]      = strategy_id
        row_data["resultado"]       = "BALANCE_INICIO_MES"
        row_data["balance_despues"] = round(balance, 4)

        col_map = {str(cell.value).lower(): cell.column for cell in ws[1] if cell.value}
        col_order = [name for name, _ in sorted(col_map.items(), key=lambda x: x[1])]
        ws.append([row_data.get(col, "") for col in col_order])

        # Colorear la fila de inicio en amarillo
        row_num = ws.max_row
        fill    = PatternFill("solid", fgColor=_COLOR_INICIO)
        font    = Font(bold=True, italic=True)
        for cell in ws[row_num]:
            cell.fill = fill
            cell.font = font
    except Exception as e:
        logger.debug(f"No se pudo agregar fila balance inicio: {e}")


def _col_map(ws) -> dict:
    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    return {str(cell.value).lower(): cell.column for cell in header_row if cell.value}


# ---------------------------------------------------------------------------
# Helpers para buscar el Excel correcto en close (puede ser mes anterior)
# ---------------------------------------------------------------------------

def _find_excel_with_order(order_id: str, strategy_id: str) -> str | None:
    """
    Busca el Excel que contiene el order_id.
    Primero el mes actual, luego el anterior.
    """
    import calendar
    now = datetime.now(timezone.utc)

    candidates = [
        get_excel_path(strategy_id, now),
    ]
    # Mes anterior
    if now.month == 1:
        prev = now.replace(year=now.year - 1, month=12, day=1)
    else:
        prev = now.replace(month=now.month - 1, day=1)
    candidates.append(get_excel_path(strategy_id, prev))

    try:
        import openpyxl
    except ImportError:
        return None

    for path in candidates:
        if not os.path.exists(path):
            continue
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            cmap = _col_map(ws)
            ocol = cmap.get("order_id")
            if not ocol:
                continue
            for row in ws.iter_rows(min_row=2):
                if str(row[ocol - 1].value or "").strip() == order_id:
                    return path
        except Exception:
            continue
    return None


# ---------------------------------------------------------------------------
# API pública
# ---------------------------------------------------------------------------

def append_excel_row(row_data: dict, strategy_id: str) -> None:
    """
    Escribe una fila de apertura (OPEN) en el Excel del mes actual.
    Crea el archivo si no existe, con fila de balance inicial si es mes nuevo.
    """
    try:
        now        = datetime.now(timezone.utc)
        excel_path = get_excel_path(strategy_id, now)
        is_new     = not os.path.exists(excel_path)
        wb, ws, _  = _get_or_create_wb(excel_path, strategy_id, is_new_month=is_new)

        # Enriquecer la fila con fecha/hora separados
        ts_str = row_data.get("timestamp_utc", "")
        try:
            ts = datetime.fromisoformat(ts_str.replace("Z", "+00:00"))
        except Exception:
            ts = now
        row_data.setdefault("fecha",    ts.strftime("%Y-%m-%d"))
        row_data.setdefault("hora_utc", ts.strftime("%H:%M:%S"))

        # Calcular notional
        margen   = float(row_data.get("margen_usdt", 0) or 0)
        leverage = float(row_data.get("leverage",    0) or 0)
        row_data.setdefault("notional_usdt", round(margen * leverage, 4) if margen and leverage else "")

        # Renombrar campos al nuevo esquema
        row_data.setdefault("precio_entrada", row_data.pop("price", ""))
        row_data.setdefault("precio_salida",  "PENDING")
        row_data.setdefault("duracion_min",   "PENDING")
        row_data.setdefault("resultado",      row_data.pop("result", "PENDING"))
        row_data.setdefault("profit_usdt",    "PENDING")
        row_data.setdefault("balance_despues","PENDING")

        cmap      = _col_map(ws)
        col_order = [name for name, _ in sorted(cmap.items(), key=lambda x: x[1])]
        ws.append([row_data.get(col, "") for col in col_order])
        wb.save(excel_path)
        logger.debug(f"Excel [{strategy_id}]: +1 fila → {Path(excel_path).name}")
    except PermissionError:
        logger.warning(f"Excel [{strategy_id}] abierto en otro programa — fila no guardada")
    except Exception as e:
        logger.error(f"Error escribiendo Excel [{strategy_id}]: {e}")


def update_excel_result(
    order_id:        str,
    result:          str,
    strategy_id:     str,
    pnl_pct:         str   = "",
    pnl_notes:       str   = "",
    precio_salida:   float = 0,
    duracion_min:    float = 0,
    profit_usdt:     float = 0,
    balance_antes:   float = 0,
    balance_despues: float = 0,
) -> bool:
    """
    Busca la fila por order_id y la actualiza con todos los datos del cierre.
    Colorea la fila según WIN/LOSS.
    """
    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl no instalado")
        return False

    excel_path = _find_excel_with_order(order_id, strategy_id)
    if not excel_path:
        logger.warning(f"Excel [{strategy_id}] order_id={order_id} no encontrado en ningún mes")
        return False

    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
    except PermissionError:
        logger.warning(f"Excel abierto [{strategy_id}] — cierra Excel para ver resultados")
        return False
    except Exception as e:
        logger.error(f"Error abriendo Excel [{strategy_id}]: {e}")
        return False

    cmap = _col_map(ws)

    def col(name):
        return cmap.get(name)

    order_col   = col("order_id")
    result_col  = col("resultado")
    if not order_col or not result_col:
        logger.error(f"Excel [{strategy_id}]: columnas 'order_id' o 'resultado' no encontradas")
        return False

    now_ts = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    found  = False

    for row in ws.iter_rows(min_row=2):
        if str(row[order_col - 1].value or "").strip() != order_id:
            continue
        current = str(row[result_col - 1].value or "").strip()
        if current in ("", "PENDING"):
            # Resultado
            row[result_col - 1].value = result

            # Precio salida
            if col("precio_salida") and precio_salida:
                row[col("precio_salida") - 1].value = precio_salida

            # Duración
            if col("duracion_min") and duracion_min:
                row[col("duracion_min") - 1].value = round(duracion_min, 1)

            # PnL %
            if col("pnl_pct") and pnl_pct:
                row[col("pnl_pct") - 1].value = pnl_pct

            # Profit USDT
            if col("profit_usdt") and profit_usdt is not None:
                row[col("profit_usdt") - 1].value = round(profit_usdt, 4)

            # Close reason — tomado de pnl_notes si no se pasa explícitamente
            if col("close_reason") and pnl_notes:
                row[col("close_reason") - 1].value = pnl_notes

            # Balance
            if col("balance_antes") and balance_antes:
                row[col("balance_antes") - 1].value = round(balance_antes, 4)
            if col("balance_despues") and balance_despues:
                row[col("balance_despues") - 1].value = round(balance_despues, 4)

            found = True

            # Colorear fila
            _color_row(ws, row[0].row, result)
        break

    if found:
        try:
            wb.save(excel_path)
            logger.info(
                f"Excel [{strategy_id}] {result} pnl={pnl_pct} "
                f"profit={profit_usdt:+.4f} USDT "
                f"balance={balance_antes:.2f}→{balance_despues:.2f} "
                f"order={order_id}"
            )
        except PermissionError:
            logger.warning(f"Excel abierto [{strategy_id}] — resultado no guardado")
            return False
        except Exception as e:
            logger.error(f"Error guardando Excel [{strategy_id}]: {e}")
            return False

    return found
